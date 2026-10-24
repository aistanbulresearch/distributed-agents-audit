"""Structural reconnaissance of the pinned upstream corpus (Phase 0, Task 3).

Run: uv run python -m audit.scripts.recon

ALLOWED: schemas, row counts, key uniqueness and join coverage, value vocabularies,
identifiers, file presence, and counts that reproduce numbers the paper already publishes.

FORBIDDEN (these are pre-registered outcomes and must not be computed here):
- flag proportions or their differences between arms or strata (Lines C, D, I);
- citation resolution or title-match rates (Line A);
- any comparison of report statistics with recomputed DE values (Line H1);
- data-consistency rates by priority tier, specificity ranks, contrast tests (H2-H4);
- reference-type composition ratios between single and distributed agents (Line F).
"""

import csv
import json
import logging
import random
import re
import statistics
from collections import Counter
from pathlib import Path

import openpyxl
import pyarrow.compute as pc
import pyarrow.parquet as pq

from audit.lib import corpus, paths, reports

logger = logging.getLogger(__name__)

OUT_DIR = paths.REPO_ROOT / "audit" / "results" / "recon"
UP = paths.UPSTREAM_ROOT
SEED = 20261023
# Paper, main text p5 and Fig. 2E: the single agent cited 11; distributed retained 711x more.
PUBLISHED_RATIO = 711
PUBLISHED_SINGLE_REFS = 11

_STAT_FAMILIES = {
    "logfc": r"log2?\s*fc|logfoldchange",
    "fdr_padj": r"\bfdr\b|p_?adj|adjusted p|q\s*[=<]",
    "nominal_p": r"\bp\s*[=<]|nominal p",
    "cell_counts": r"\bcells?\b|nuclei|n_pert|n_ctrl",
    "deg_counts": r"\bdegs?\b",
    "correlation": r"\bpearson\b|\bspearman\b|\brho\b|\br\s*=",
    "cosine": r"cosine",
    "rank": r"\brank",
}


def _csv_rows(path: Path) -> list[dict[str, str]]:
    return list(corpus.read_csv_checked(path))


def ledger_inventory(release: Path) -> dict:
    out = {}
    for name in ("findings", "references", "evidence"):
        rows = _csv_rows(release / "ledgers" / f"{name}.csv")
        out[name] = {"rows": len(rows), "columns": list(rows[0].keys())}
    return out


def classified_inventory(path: Path, findings: list[corpus.Finding]) -> dict:
    rows = _csv_rows(path)
    cols = list(rows[0].keys())
    combo = Counter(
        (r.get("arm", ""), r.get("set_role", ""), r.get("valid_for_primary", "")) for r in rows
    )
    ledger_keys = {(f.target_gene, f.finding_id) for f in findings}
    keyed = [r for r in rows if r.get("finding_id") and r.get("report_gene_target")]
    linked = sum((r["report_gene_target"], r["finding_id"]) in ledger_keys for r in keyed)
    return {
        "rows": len(rows),
        "columns": cols,
        "rows_by_arm_setrole_validprimary": {" | ".join(k): v for k, v in sorted(combo.items())},
        "rows_with_finding_key": len(keyed),
        "finding_keys_found_in_ledger": linked,
    }


def control_inventory(all_findings: list[corpus.Finding], screen: set[str]) -> dict:
    """``screen``: perturbation names from Table 3 (depletion sheet), the full screen list."""
    targets = {f.target_gene for f in all_findings}
    controls = sorted(t for t in targets if corpus.is_control(t))
    non_control = {t for t in targets if not corpus.is_control(t)}
    return {
        "distinct_targets": len(targets),
        "control_targets": len(controls),
        "control_examples": controls[:5],
        "control_findings": sum(corpus.is_control(f.target_gene) for f in all_findings),
        "screen_perturbations_table3": len(screen),
        "ledger_targets_not_in_table3": sorted(non_control - screen)[:20],
        "table3_perturbations_not_in_ledger": sorted(screen - non_control)[:20],
    }


def single_agent_refs(scaling_dir: Path) -> dict:
    out = {}
    for seed_dir in sorted(scaling_dir.glob("seed-*")):
        path = seed_dir / "n0893" / "q3_references.jsonl"
        rows = list(reports.iter_jsonl(path))
        out[seed_dir.name] = {
            "rows": len(rows),
            "keys": sorted(rows[0].keys()) if rows else [],
            "source_types": dict(Counter(r.get("source_type", "") for r in rows)),
        }
    return out


def scaling_targets(media3: Path) -> dict[str, list[str]]:
    """Per-seed target order from Table 2 (one column per seed, header ``seed_NN``)."""
    ws = openpyxl.load_workbook(media3, read_only=True)["Scaling Perturbation Splits"]
    rows = list(ws.iter_rows(values_only=True))
    return {
        str(seed): [str(r[i]) for r in rows[1:] if r[i]]
        for i, seed in enumerate(rows[0])
    }


def distributed_comparator(refs: dict, split: dict[str, list[str]]) -> dict:
    sets = {seed: set(targets) for seed, targets in split.items()}
    wanted = set().union(*sets.values())
    rows = [r for (t, _), r in refs.items() if t in wanted]
    http = [r for r in rows if r.url.strip().lower().startswith(("http://", "https://"))]
    return {
        "targets_per_seed": {seed: len(s) for seed, s in sets.items()},
        "seeds_share_one_target_set": len({frozenset(s) for s in sets.values()}) == 1,
        "targets_union": len(wanted),
        "union_targets_with_release_report": len({t for t, _ in refs} & wanted),
        "reference_rows": len(rows),
        "rows_with_http_url": len(http),
        "implied_by_published_711x_of_11": PUBLISHED_RATIO * PUBLISHED_SINGLE_REFS,
    }


def literature_comparisons(report_path: Path) -> dict:
    n = with_block = n_rows = n_ctx = n_invalid = 0
    for rec in reports.iter_jsonl(report_path):
        n += 1
        block = reports.extract_jsonl_block(rec["final_report"], "Literature Comparisons JSONL")
        if block is None:
            continue
        with_block += 1
        n_rows += len(block.rows)
        n_ctx += sum("context_match" in row for row in block.rows)
        n_invalid += block.n_invalid
    return {"reports": n, "reports_with_block": with_block, "rows": n_rows,
            "rows_with_context_match": n_ctx, "invalid_lines": n_invalid}


def _report_map(path: Path, text_key: str) -> dict[str, str]:
    return {rec["gene_target"]: rec[text_key] for rec in reports.iter_jsonl(path)}


def snapshot_comparison(release: Path) -> dict:
    rel = _report_map(release / "reports" / "reports.jsonl", "final_report")
    out: dict[str, object] = {"release_reports": len(rel)}
    for name in ("june", "july"):
        snap = _report_map(UP / "data" / "runs" / "reports" / name / "reports.jsonl", "report")
        shared = rel.keys() & snap.keys()
        out[name] = {
            "reports": len(snap),
            "shared_targets": len(shared),
            "only_in_snapshot": len(snap.keys() - rel.keys()),
            "only_in_release": len(rel.keys() - snap.keys()),
            "identical_text": sum(rel[t] == snap[t] for t in shared),
        }
    return out


def url_vocabulary(refs: dict) -> dict:
    hosts = Counter(reports.url_host(r.url) for r in refs.values())
    return {"non_http_or_empty": hosts.pop("", 0), "distinct_hosts": len(hosts),
            "top_hosts": dict(hosts.most_common(30))}


def build_module_status() -> dict:
    gitignore = (UP / ".gitignore").read_text(encoding="utf-8").splitlines()
    return {"build_dir_exists": (UP / "src" / "distributed_agents" / "build").exists(),
            "gitignore_build_lines": [line for line in gitignore if "build" in line]}


def source_policy_tasks() -> dict:
    out = {}
    for toml in sorted((UP / "src" / "bioeval" / "benchmarks").rglob("*.toml")):
        text = toml.read_text(encoding="utf-8")
        mode = re.search(r'^\s*mode\s*=\s*"([^"]+)"', text, re.M)
        has_policy = "[source_policy]" in text and mode is not None
        out[toml.relative_to(UP).as_posix()] = mode.group(1) if has_policy and mode else None
    return out


def de_structure(path: Path) -> dict:
    table = pq.read_table(path)
    pairs = Counter(zip(table["gene_target"].to_pylist(), table["group_name"].to_pylist(),
                        strict=True))
    sizes = sorted(pairs.values())
    padj = table["pvals_adj"]
    groups = sorted(set(table["group_name"].to_pylist()))
    tsc = {}
    for target in ("Tsc1", "Tsc2"):
        for group in (g for g in groups if "CB" in g and "GABA" in g.upper()):
            mask = pc.and_(pc.equal(table["gene_target"], target),
                           pc.equal(table["group_name"], group))
            sub = table.filter(mask)
            if sub.num_rows:
                tsc[f"{target} / {group}"] = {
                    "rows": sub.num_rows,
                    "n_pert_matched": sorted(set(sub["n_pert_matched"].to_pylist())),
                    "n_ctrl_matched": sorted(set(sub["n_ctrl_matched"].to_pylist())),
                }
    return {
        "rows": table.num_rows,
        "distinct_targets": len(set(table["gene_target"].to_pylist())),
        "distinct_groups": len(groups),
        "target_group_pairs": len(pairs),
        "rows_per_pair_min_median_max": [sizes[0], statistics.median(sizes), sizes[-1]],
        "pvals_adj_min_max": [pc.min(padj).as_py(), pc.max(padj).as_py()],
        "rows_with_pvals_adj_ge_0.1": pc.sum(pc.greater_equal(padj, 0.1)).as_py(),
        "groups": groups,
        "tsc_cb_gaba": tsc,
    }


def evidence_vocabulary(release: Path) -> dict:
    rows = _csv_rows(release / "ledgers" / "evidence.csv")
    stats = [r["statistics"] for r in rows]
    families = {name: sum(bool(re.search(rx, s, re.I)) for s in stats)
                for name, rx in _STAT_FAMILIES.items()}
    examples = random.Random(SEED).sample(stats, 6)
    return {"rows": len(stats), "empty": sum(not s.strip() for s in stats),
            "rows_without_cell_type": sum(not r["cell_type"].strip() for r in rows),
            "rows_matching_family": families, "examples": examples}


def released_tables() -> dict:
    out = {}
    for i in range(2, 7):
        wb = openpyxl.load_workbook(paths.PREPRINT_SUPP_DIR / f"media-{i}.xlsx", read_only=True)
        out[f"media-{i}"] = {
            ws.title: {"rows": ws.max_row,
                       "header": [str(c) for c in next(ws.iter_rows(values_only=True))]}
            for ws in wb.worksheets
        }
    return out


def survey_inventory() -> dict:
    base = UP / "data" / "inputs" / "survey"
    findings = list(csv.DictReader(open(base / "findings.csv", encoding="utf-8")))
    comparisons = list(csv.DictReader(open(base / "comparisons.csv", encoding="utf-8")))
    return {"findings_rows": len(findings),
            "findings_by_source": dict(Counter(r["source"] for r in findings)),
            "comparisons_rows": len(comparisons),
            "comparison_columns": list(comparisons[0].keys())}


def _sheet_rows(book: str, sheet: str) -> list[dict]:
    ws = openpyxl.load_workbook(paths.PREPRINT_SUPP_DIR / book, read_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    return [dict(zip(header, r, strict=False)) for r in rows[1:] if any(v is not None for v in r)]


def _ledger_link(rows: list[dict], gene_col: str, id_col: str, ledger_keys: set) -> dict:
    keys = [(str(r[gene_col]), str(r[id_col])) for r in rows]
    return {
        "finding_id_examples": [finding_id for _, finding_id in keys[:3]],
        "keys_found_in_ledger": sum(key in ledger_keys for key in keys),
    }


def heatmap_inventory(ledger_keys: set) -> dict:
    """Table 2 'Heatmap Finding Flags': counts by confidence only (flag split is Line D)."""
    rows = _sheet_rows("media-3.xlsx", "Heatmap Finding Flags")
    return {
        "rows": len(rows),
        "rows_by_confidence": dict(Counter(str(r["confidence"]) for r in rows)),
        "distinct_perturbations": len({r["perturbation"] for r in rows}),
        **_ledger_link(rows, "perturbation", "finding_id", ledger_keys),
    }


def priority_table_inventory(ledger_keys: set) -> dict:
    rows = _sheet_rows("media-5.xlsx", "analysis_table")
    return {"rows": len(rows), "distinct_genes": len({r["gene"] for r in rows}),
            **_ledger_link(rows, "gene", "finding_id", ledger_keys)}


def depletion_perturbations() -> set[str]:
    return {str(r["Perturbation"]) for r in _sheet_rows("media-4.xlsx",
                                                         "depletion rank and predictions")}


def runtime_runs() -> dict:
    rows = _sheet_rows("media-3.xlsx", "Runtime Summary Statistics")
    return {"runs": sorted({str(r["run"]) for r in rows}),
            "metrics": sorted({str(r["metric"]) for r in rows})}


def survey_pair_texts() -> dict:
    """Whether each pair's two titles map to a full text in the survey findings table."""
    base = UP / "data" / "inputs" / "survey"
    findings = list(csv.DictReader(open(base / "findings.csv", encoding="utf-8")))
    titles = {r["finding_title"].strip() for r in findings}
    pairs = list(csv.DictReader(open(base / "comparisons.csv", encoding="utf-8")))
    report_144 = (UP / "data" / "runs" / "survey" / "distributed_agents" / "report_all_144.md")
    text_144 = report_144.read_text(encoding="utf-8")
    return {
        "pairs": len(pairs),
        "codex_title_has_full_text": sum(p["codex_title"].strip() in titles for p in pairs),
        "distributed_title_has_full_text": sum(
            p["distributed_agents_title"].strip() in titles for p in pairs),
        "distributed_title_in_report_all_144": sum(
            p["distributed_agents_title"].strip() in text_144 for p in pairs),
    }


def downstream_parquet(path: Path) -> dict:
    table = pq.read_table(path, columns=["gene_target", "cell_type", "pvals_adj"])
    padj = table["pvals_adj"]
    return {"rows": table.num_rows,
            "distinct_targets": len(set(table["gene_target"].to_pylist())),
            "distinct_cell_types": len(set(table["cell_type"].to_pylist())),
            "pvals_adj_max": pc.max(padj).as_py()}


def holdout_run_files() -> dict:
    files = Counter(p.name for p in (UP / "data" / "runs" / "holdout").rglob("*") if p.is_file())
    return dict(files.most_common())


def other_reference_files() -> list[str]:
    hits = [p for p in (UP / "data").rglob("*reference*") if p.is_file()]
    return sorted(p.relative_to(UP).as_posix() for p in hits if "scaling" not in p.parts)[:40]


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    paths.check_upstream_pin()
    release = paths.release_dir()
    all_findings = corpus.load_findings(release, include_controls=True)
    refs = corpus.load_references(release)
    de_path = UP / "data" / "inputs" / "biology" / "differential_expression.parquet"
    split = scaling_targets(paths.PREPRINT_SUPP_DIR / "media-3.xlsx")
    ledger_keys = {(f.target_gene, f.finding_id) for f in all_findings}

    inventory = {
        "ledgers": ledger_inventory(release),
        "classified_findings": classified_inventory(
            UP / "data" / "inputs" / "literature" / "final_classified_findings.csv", all_findings),
        "heatmap_finding_flags_table2": heatmap_inventory(ledger_keys),
        "priority_scores_table4": priority_table_inventory(ledger_keys),
        "runtime_summary_table2": runtime_runs(),
        "survey_pair_texts": survey_pair_texts(),
        "downstream_fdr_0_1_parquet": downstream_parquet(
            UP / "data" / "inputs" / "scaling" / "downstream_fdr_0_1.parquet"),
        "holdout_run_files": holdout_run_files(),
        "controls": control_inventory(all_findings, depletion_perturbations()),
        "single_agent_refs_n0893": single_agent_refs(UP / "data" / "runs" / "scaling"),
        "distributed_comparator_893": distributed_comparator(refs, split),
        "literature_comparisons": literature_comparisons(release / "reports" / "reports.jsonl"),
        "snapshots": snapshot_comparison(release),
        "url_vocabulary": url_vocabulary(refs),
        "build_module": build_module_status(),
        "source_policy_tasks": source_policy_tasks(),
        "differential_expression": de_structure(de_path),
        "evidence_statistics": evidence_vocabulary(release),
        "released_tables": released_tables(),
        "survey": survey_inventory(),
        "other_reference_files": other_reference_files(),
    }
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    out = OUT_DIR / "inventory.json"
    out.write_text(json.dumps(inventory, indent=2, ensure_ascii=False) + "\n",
                   encoding="utf-8", newline="\n")
    logger.info("wrote %s", out)


if __name__ == "__main__":
    main()
